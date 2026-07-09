import streamlit as st
import pdfplumber
import requests

from docx import Document
import pandas as pd

# ==================================================
# CONFIGURAÇÃO
# ==================================================

st.set_page_config(
    page_title="Clipping de Diários",
    layout="wide"
)

PALAVRAS = [
    "SEPPEN",
    "PENA JUSTA",
    "CEPPRJ",
    "GMF",
    "superlotação prisional"
]

# ==================================================
# FUNÇÕES
# ==================================================

def baixar_pdf(url, nome_arquivo):

    resposta = requests.get(url, timeout=60)

    resposta.raise_for_status()

    with open(nome_arquivo, "wb") as f:
        f.write(resposta.content)

    return nome arquivo

def extrair_texto(pdf_path):

    texto = ""

    with pdfplumber.open(pdf_path) as pdf:

        for pagina in pdf.pages:

            txt = pagina.extract_text()

            if txt:
                texto += txt + "\n"

    return texto


def obter_linhas(texto):

    linhas = []

    for linha in texto.split("\n"):

        linha = linha.strip()

        if linha:
            linhas.append(linha)

    return linhas


def localizar_ocorrencias(linhas, palavras):

    resultados = []

    for i, linha in enumerate(linhas):

        for palavra in palavras:

            if palavra.lower() in linha.lower():

                resultados.append({
                    "linha": i,
                    "palavra": palavra,
                    "texto": linha
                })

    return resultados


def gerar_word(blocos):

    doc = Document()

    doc.add_heading(
        "CLIPPING DIÁRIO",
        level=0
    )

    for i, bloco in enumerate(
    blocos,
    start=1
):

    doc.add_heading(
        f"Ocorrência {i}",
        level=1
    )

    doc.add_paragraph(
        f"Documento: {bloco['documento']}"
    )

    doc.add_paragraph(
        f"Palavra-chave: {bloco['palavra']}"
    )

    doc.add_paragraph(
        bloco["conteudo"]
    )

        doc.add_paragraph(bloco)

    caminho = "clipping.docx"

    doc.save(caminho)

    return caminho


# ==================================================
# SESSION STATE
# ==================================================

if "inicio" not in st.session_state:
    st.session_state.inicio = None

if "fim" not in st.session_state:
    st.session_state.fim = None

if "clipping" not in st.session_state:
    st.session_state.clipping = []

# ==================================================
# TÍTULO
# ==================================================

st.title("📑 Clipping de Diários Oficiais")

st.subheader("URLs dos Diários")

col1, col2 = st.columns(2)

with col1:

    url_cnj = st.text_input(
        "CNJ"
    )

    url_cnmp = st.text_input(
        "CNMP"
    )

with col2:

    url_tjrj = st.text_input(
        "TJRJ"
    )

    url_ioerj = st.text_input(
        "IOERJ"
    )

processar = st.button(
    "Processar Diários"
)

# ==================================================
# PROCESSAMENTO
# ==================================================

if processar:

documentos = [
    ("CNJ", url_cnj),
    ("CNMP", url_cnmp),
    ("TJRJ", url_tjrj),
    ("IOERJ", url_ioerj)
]

    with open("temp.pdf", "wb") as f:
        f.write(arquivo.read())

    for nome_doc, url in documentos:

    if not url.strip():
        continue

    try:

        caminho_pdf = f"{nome_doc}.pdf"

        baixar_pdf(
            url,
            caminho_pdf
        )

        texto = extrair_texto(
            caminho_pdf
        )

        linhas = obter_linhas(
            texto
        )

        ocorrencias = localizar_ocorrencias(
            linhas,
            PALAVRAS
        )

        st.header(nome_doc)

        st.success(
            f"{len(ocorrencias)} ocorrência(s)"
        )

    # ==========================================
    # NENHUMA OCORRÊNCIA
    # ==========================================

    if len(ocorrencias) == 0:

        st.warning(
            "Nenhuma palavra-chave localizada."
        )

    # ==========================================
    # OCORRÊNCIAS
    # ==========================================

    else:

        st.subheader("Ocorrências Encontradas")

        for indice, item in enumerate(ocorrencias):

            with st.expander(
                f"{item['palavra']} - Linha {item['linha']}"
            ):

                linha_chave = item["linha"]

                inicio_contexto = max(
                    0,
                    linha_chave - 30
                )

                fim_contexto = min(
                    len(linhas),
                    linha_chave + 70
                )

                st.info(
                    f"Texto localizado: {item['texto']}"
                )

                for i in range(
                    inicio_contexto,
                    fim_contexto
                ):

                    col1, col2, col3 = st.columns(
                        [1, 1, 10]
                    )

                    with col1:

                        if st.button(
                            "Início",
                            key=f"inicio_{indice}_{i}"
                        ):
                            st.session_state.inicio = i

                    with col2:

                        if st.button(
                            "Fim",
                            key=f"fim_{indice}_{i}"
                        ):
                            st.session_state.fim = i

                    with col3:

                        texto_linha = linhas[i]

                        if (
                            item["palavra"].lower()
                            in texto_linha.lower()
                        ):

                            st.markdown(
                                f"🟨 **{i} - {texto_linha}**"
                            )

                        else:

                            st.write(
                                f"{i} - {texto_linha}"
                            )

        st.divider()

        # ==========================================
        # CONTROLE DA SELEÇÃO
        # ==========================================

        col_a, col_b, col_c = st.columns(3)

        with col_a:
            st.success(
                f"Início: {st.session_state.inicio}"
            )

        with col_b:
            st.success(
                f"Fim: {st.session_state.fim}"
            )

        with col_c:

            if st.button(
                "Limpar Seleção"
            ):

                st.session_state.inicio = None
                st.session_state.fim = None

                st.rerun()

        trecho = ""

        # ==========================================
        # TRECHO SELECIONADO
        # ==========================================

        if (
            st.session_state.inicio is not None
            and
            st.session_state.fim is not None
        ):

            inicio = min(
                st.session_state.inicio,
                st.session_state.fim
            )

            fim = max(
                st.session_state.inicio,
                st.session_state.fim
            )

            trecho = "\n".join(
                linhas[inicio:fim + 1]
            )

            st.subheader("Prévia")

            st.text_area(
                "Trecho selecionado",
                trecho,
                height=300
            )

            if st.button(
                "Adicionar ao Clipping"
            ):

                st.session_state.clipping.append({

    "documento": nome_doc,

    "palavra": item["palavra"],

    "conteudo": trecho

})

                st.success(
                    "Trecho adicionado."
                )

        # ==========================================
        # CLIPPING ACUMULADO
        # ==========================================

        st.divider()

        st.subheader(
            "Clipping Atual"
        )

        if len(st.session_state.clipping) == 0:

            st.info(
                "Nenhum trecho selecionado."
            )

        else:

            for i, bloco in enumerate(
                st.session_state.clipping,
                start=1
            ):

                with st.expander(
                    f"Trecho {i}"
                ):

                    st.write(
    f"Documento: {bloco['documento']}"
)

st.write(
    f"Palavra: {bloco['palavra']}"
)

st.text(
    bloco["conteudo"]
)

        # ==========================================
        # GERAR WORD
        # ==========================================

        if (
            len(
                st.session_state.clipping
            ) > 0
        ):

            if st.button(
                "Gerar Word"
            ):

                arquivo_word = gerar_word(
                    st.session_state.clipping
                )

                st.success(
                    "Documento criado."
                )

                with open(
                    arquivo_word,
                    "rb"
                ) as f:

                    st.download_button(
                        "📥 Baixar Clipping",
                        data=f,
                        file_name="clipping.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
